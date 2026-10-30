from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def build_workbook(exam, rows, review) -> bytes:
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="4A2C8F")

    def sheet(ws, header, data):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in data:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))

    ws = wb.active
    ws.title = "Summary"
    sheet(ws, ["Roll number", "Name", "Marks obtained", "Total marks", "Percentage", "Correct", "Wrong", "Unanswered"],
          [[r["roll_no"], r["name"], r["total"], exam.total_marks, r["pct"], r["correct"], r["wrong"], r["unanswered"]] for r in rows])
    sheet(wb.create_sheet("Detail"), ["Roll number", "Question", "Student answer", "Correct answer", "Result", "Marks"],
          [[r["roll_no"], a["number"], a["student_answer"], a["correct_answer"], a["result"], a["marks"]] for r in rows for a in r["answers"]])
    sheet(wb.create_sheet("Unreadable sheets"), ["Roll number", "File", "Unreadable", "Reason"],
          [[u["roll_no"] or "(unreadable)", u["filename"], u["detail"], u["reason"]] for u in review])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
