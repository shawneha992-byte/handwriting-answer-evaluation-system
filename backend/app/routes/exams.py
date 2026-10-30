import csv
import io
import re
import uuid
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Response, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.models import Exam, Question, Submission
from app.services import llm_service
from app.services.export import build_workbook
from app.services.processing import MIME, run_batch

router = APIRouter(prefix="/exams", tags=["exams"])
FILE_TYPES = {".pdf", ".jpg", ".jpeg", ".png"}


# ---------------------------------------------------------------- helpers
def read_upload(f: UploadFile, allowed: set[str]) -> tuple[bytes, str]:
    suffix = Path(f.filename or "").suffix.lower()
    if suffix not in allowed:
        raise HTTPException(400, f"{f.filename}: unsupported file type. Allowed: {', '.join(sorted(allowed))}")
    data = f.file.read()
    if len(data) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(400, f"{f.filename}: larger than {settings.MAX_UPLOAD_MB} MB")
    return data, suffix


def get_exam(db: Session, exam_id: int) -> Exam:
    exam = db.get(Exam, exam_id)
    if not exam:
        raise HTTPException(404, "Exam not found")
    return exam


def key_from_table(rows) -> dict:
    """Read question-number -> answer pairs from the first two columns of a CSV/Excel table."""
    out = {}
    for row in rows:
        if len(row) < 2 or row[0] is None:
            continue
        m = re.match(r"\s*Q?\s*(\d+)", str(row[0]), re.I)
        if m and row[1] is not None and str(row[1]).strip() != "":
            val = row[1]
            out[m.group(1)] = str(int(val)) if isinstance(val, float) and val == int(val) else str(val).strip()
    return out


def results_for(exam: Exam) -> list[dict]:
    rows = []
    for s in exam.submissions:
        if s.status != "done":
            continue
        ans = [{"number": a.question_number, "student_answer": a.student_answer, "correct_answer": a.correct_answer,
                "result": a.result, "marks": a.marks} for a in s.answers]
        rows.append({
            "id": s.id, "roll_no": s.roll_no or "-", "name": s.student_name or "-", "total": s.total_score or 0,
            "pct": round((s.total_score or 0) / exam.total_marks * 100) if exam.total_marks else 0,
            "correct": sum(a["result"] == "correct" for a in ans), "wrong": sum(a["result"] == "wrong" for a in ans),
            "unanswered": sum(a["result"] == "unanswered" for a in ans), "manual": s.manual, "answers": ans,
        })
    return sorted(rows, key=lambda r: str(r["roll_no"]))


def review_for(exam: Exam) -> list[dict]:
    return [{"id": s.id, "roll_no": s.roll_no, "filename": s.filename, "name": s.student_name, "detail": s.review_detail or "",
             "reason": s.review_reason or ""} for s in exam.submissions if s.status == "needs_review"]


# ---------------------------------------------------------------- step 1-3: read paper and key (nothing saved yet)
@router.post("/parse-paper")
def parse_paper(file: UploadFile = File(...)):
    data, suffix = read_upload(file, FILE_TYPES)
    try:
        return llm_service.parse_paper(data, MIME[suffix])
    except llm_service.LLMError as exc:
        raise HTTPException(502, str(exc))


@router.post("/parse-key")
def parse_key(file: UploadFile = File(...)):
    data, suffix = read_upload(file, FILE_TYPES | {".csv", ".xlsx"})
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        answers = key_from_table(csv.reader(io.StringIO(text)))
    elif suffix == ".xlsx":
        ws = load_workbook(io.BytesIO(data), data_only=True).active
        answers = key_from_table(ws.iter_rows(values_only=True))
    else:
        try:
            answers = llm_service.parse_key_image(data, MIME[suffix]).get("answers", {})
        except llm_service.LLMError as exc:
            raise HTTPException(502, str(exc))
    if not answers:
        raise HTTPException(422, "No answers found. Use two columns: question number, correct answer.")
    return {"answers": {str(k): str(v) for k, v in answers.items()}}


# ---------------------------------------------------------------- create exam and upload sheets
class QuestionIn(BaseModel):
    number: int
    text: str = ""
    type: str = "MCQ"
    marks: float
    correct_answer: str = ""


class ExamIn(BaseModel):
    title: str
    negative_marking: float = 0
    questions: list[QuestionIn]


@router.post("")
def create_exam(body: ExamIn, db: Session = Depends(get_db)):
    if not body.questions:
        raise HTTPException(422, "The exam has no questions")
    missing = [f"Q{q.number}" for q in body.questions if not q.correct_answer.strip()]
    if missing:
        raise HTTPException(422, f"Answer key is missing for: {', '.join(missing)}")
    exam = Exam(title=body.title.strip() or "Untitled exam", negative_marking=min(0, body.negative_marking),
                total_marks=sum(q.marks for q in body.questions))
    exam.questions = [Question(number=q.number, text=q.text, type=q.type, marks=q.marks, correct_answer=q.correct_answer.strip())
                      for q in body.questions]
    db.add(exam)
    db.commit()
    return {"id": exam.id}


@router.post("/{exam_id}/sheets")
def upload_sheets(exam_id: int, background: BackgroundTasks, files: list[UploadFile] = File(...), db: Session = Depends(get_db)):
    exam = get_exam(db, exam_id)
    folder = settings.UPLOAD_DIR / str(exam.id)
    folder.mkdir(parents=True, exist_ok=True)
    subs = []
    for f in files:
        data, suffix = read_upload(f, FILE_TYPES)
        path = folder / f"{uuid.uuid4().hex}{suffix}"  # generated name: no overwrites, no path tricks
        path.write_bytes(data)
        subs.append(Submission(exam_id=exam.id, filename=Path(f.filename).name, stored_path=str(path)))
    db.add_all(subs)
    db.commit()
    background.add_task(run_batch, [s.id for s in subs])
    return {"queued": len(subs)}


# ---------------------------------------------------------------- read results
@router.get("")
def list_exams(db: Session = Depends(get_db)):
    exams = db.query(Exam).order_by(Exam.id.desc()).all()
    return [{"id": e.id, "title": e.title, "sheets": len(e.submissions)} for e in exams]


@router.get("/{exam_id}/status")
def status(exam_id: int, db: Session = Depends(get_db)):
    exam = get_exam(db, exam_id)
    return [{"id": s.id, "filename": s.filename, "roll_no": s.roll_no, "status": s.status} for s in exam.submissions]


@router.get("/{exam_id}/results")
def results(exam_id: int, db: Session = Depends(get_db)):
    exam = get_exam(db, exam_id)
    rows = results_for(exam)
    pcts = [r["pct"] for r in rows]
    return {
        "exam": {"id": exam.id, "title": exam.title, "total_marks": exam.total_marks, "negative_marking": exam.negative_marking},
        "stats": {"students": len(rows), "average": round(sum(pcts) / len(pcts)) if pcts else 0,
                  "highest": max(pcts, default=0), "lowest": min(pcts, default=0), "needs_review": len(review_for(exam))},
        "rows": rows,
    }


@router.get("/{exam_id}/unreadable")
def unreadable(exam_id: int, db: Session = Depends(get_db)):
    return review_for(get_exam(db, exam_id))


@router.get("/{exam_id}/export")
def export(exam_id: int, db: Session = Depends(get_db)):
    exam = get_exam(db, exam_id)
    blob = build_workbook(exam, results_for(exam), review_for(exam))
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", exam.title) or "results"
    return Response(blob, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}_results.xlsx"'})


# ---------------------------------------------------------------- teacher enters marks for an unreadable sheet
class ManualIn(BaseModel):
    marks: float


@router.post("/submissions/{sub_id}/manual")
def manual_marks(sub_id: int, body: ManualIn, db: Session = Depends(get_db)):
    sub = db.get(Submission, sub_id)
    if not sub:
        raise HTTPException(404, "Sheet not found")
    total = sub.exam.total_marks
    if not 0 <= body.marks <= total:
        raise HTTPException(422, f"Marks must be between 0 and {total}")
    sub.total_score, sub.manual, sub.status = body.marks, True, "done"
    db.commit()
    return {"ok": True}
